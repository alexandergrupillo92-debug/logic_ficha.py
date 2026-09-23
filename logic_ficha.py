import os
from lxml import etree
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor

try:
    from PIL import Image as PILImage, ImageOps
except ImportError:
    PILImage = None
    ImageOps = None

try:
    _RESAMPLE = PILImage.Resampling.LANCZOS
except AttributeError:
    _RESAMPLE = PILImage.LANCZOS if PILImage else None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CARPETA_TEMP_FOTOS = os.path.join(BASE_DIR, "data", "_temp_fotos")
RUTA_MEMBRETE = os.path.join(BASE_DIR, "assets", "membrete_fondo.jpg")

LADO_PX_FOTO = 700  # resolución máxima por lado; de sobra para imprimir a 8x8 cm

# Colores (mismos del PDF)
AZUL = "1A365D"
GRIS_ETIQUETA = "E2E8F0"
GRIS_FILA = "F7FAFC"
GRIS_BORDE = "CBD5E0"

# Ancho útil de la hoja: carta (612 pt) - márgenes de 40 pt por lado
MARGEN_LATERAL_PT = 40
MARGEN_SUPERIOR_PT = 110   # no pisar el logo superior del membrete
MARGEN_INFERIOR_PT = 80    # no pisar el pie del membrete

TIPO_LETRA = "Arial"
TIPO_LETRA_MEMBRETE = "Arial Narrow"


# ---------------------------------------------------------------------------
# FOTOS
# ---------------------------------------------------------------------------
def recortar_cuadrado(ruta_foto):
    """Recorta la foto al centro (1:1), corrige la rotación EXIF y la reduce a
    un tamaño moderado para que el Word no pese de más. Devuelve la ruta de la
    foto optimizada, o la original si algo falla."""
    if PILImage is None:
        return ruta_foto
    try:
        os.makedirs(CARPETA_TEMP_FOTOS, exist_ok=True)
        im = PILImage.open(ruta_foto)
        if ImageOps is not None:
            im = ImageOps.exif_transpose(im)
        im = im.convert("RGB")
        w, h = im.size
        lado = min(w, h)
        izq = (w - lado) // 2
        arriba = (h - lado) // 2
        recorte = im.crop((izq, arriba, izq + lado, arriba + lado))
        if lado > LADO_PX_FOTO:
            recorte = recorte.resize((LADO_PX_FOTO, LADO_PX_FOTO), _RESAMPLE)
        nombre = f"cuad_{os.path.splitext(os.path.basename(ruta_foto))[0]}.jpg"
        ruta_salida = os.path.join(CARPETA_TEMP_FOTOS, nombre)
        recorte.save(ruta_salida, "JPEG", quality=78, optimize=True)
        return ruta_salida
    except Exception:
        return ruta_foto


# ---------------------------------------------------------------------------
# UTILIDADES DE FORMATO (python-docx queda corto para tablas, se completa con XML)
# ---------------------------------------------------------------------------
def _rgb(hex_color):
    return RGBColor.from_string(hex_color)


def _fuente(run, size, bold=False, color=None, nombre=TIPO_LETRA):
    run.font.name = nombre
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), nombre)


def _formato_parrafo(p, align=None, antes=0, despues=0, interlineado=None,
                     mantener_con_siguiente=False, salto_antes=False):
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    if interlineado:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(interlineado)
    if align is not None:
        p.alignment = align
    pf.keep_with_next = mantener_con_siguiente
    pf.page_break_before = salto_antes


def _escribir(p, texto, size=8.5, bold=False, color=None):
    """Escribe texto respetando saltos de línea."""
    lineas = (texto or "").replace("\r", "").split("\n")
    run = p.add_run()
    _fuente(run, size, bold, color)
    for i, linea in enumerate(lineas):
        run.add_text(linea)
        if i < len(lineas) - 1:
            run.add_break()
    return run


def _sombrear_celda(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    for viejo in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(viejo)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.insert_element_before(
        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark"
    )


def _borde(nombre, sz, color):
    e = OxmlElement(f"w:{nombre}")
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "0")
    e.set(qn("w:color"), color)
    return e


def _nueva_tabla(doc, filas, anchos_pt, borde_exterior=True, rejilla=True,
                 pad_v=3, pad_h=6):
    """Tabla de ancho fijo, con bordes y márgenes internos como en el PDF."""
    tabla = doc.add_table(rows=filas, cols=len(anchos_pt))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    tbl_pr = tabla._tbl.tblPr

    # Ancho total
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert_element_before(tbl_w, "w:jc", "w:tblLayout", "w:tblLook")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(anchos_pt) * 20)))

    # Bordes
    bordes = OxmlElement("w:tblBorders")
    if borde_exterior:
        for lado in ("top", "left", "bottom", "right"):
            bordes.append(_borde(lado, 6, AZUL))       # 0.75 pt
    if rejilla:
        bordes.append(_borde("insideH", 4, GRIS_BORDE))  # 0.5 pt
        bordes.append(_borde("insideV", 4, GRIS_BORDE))
    if len(bordes):
        tbl_pr.insert_element_before(
            bordes, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
            "w:tblCaption", "w:tblDescription"
        )

    # Márgenes internos de celda
    mar = OxmlElement("w:tblCellMar")
    for lado, valor in (("top", pad_v), ("left", pad_h), ("bottom", pad_v), ("right", pad_h)):
        m = OxmlElement(f"w:{lado}")
        m.set(qn("w:w"), str(int(valor * 20)))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tbl_pr.insert_element_before(mar, "w:tblLook", "w:tblCaption", "w:tblDescription")

    # Anchos de columna y de cada celda
    for i, ancho in enumerate(anchos_pt):
        tabla.columns[i].width = Pt(ancho)
    for fila in tabla.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = Pt(anchos_pt[i])
    return tabla


def _fila_sin_partir(fila):
    tr_pr = fila._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def _texto_en_celda(celda, texto, size=8.5, bold=False, color=None,
                    align=None, interlineado=11):
    p = celda.paragraphs[0]
    _formato_parrafo(p, align=align, interlineado=interlineado)
    _escribir(p, texto, size, bold, color)


# ---------------------------------------------------------------------------
# MEMBRETE (imagen de fondo en el encabezado + texto de las direcciones)
# ---------------------------------------------------------------------------
def _lineas_membrete(texto):
    """Devuelve las líneas del texto de direcciones. La opción completa (la más
    larga) se parte en 2 líneas para que quepa bajo la línea de la Alcaldía."""
    t = " ".join((texto or "").split())
    t = t.replace("SERVICIOS-", "SERVICIOS -").replace("MUNICIPAL-", "MUNICIPAL -")
    if len(t) <= 60:
        return [t] if t else []
    partes = [x.strip() for x in t.split(" - ")]
    if len(partes) == 1:
        return [t]
    return [partes[0], " - ".join(partes[1:])]


def _fondo_en_encabezado(doc, parrafo, ruta_imagen):
    """Inserta la imagen del membrete a página completa, detrás del texto,
    anclada a la página (igual que en tus informes)."""
    run = parrafo.add_run()
    run.add_picture(ruta_imagen, width=Inches(8.5), height=Inches(11))
    inline = run._r.xpath(".//wp:inline")[0]
    wp = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

    anchor = etree.Element(f"{{{wp}}}anchor")
    for k, v in dict(distT="0", distB="0", distL="0", distR="0", simplePos="0",
                     relativeHeight="1", behindDoc="1", locked="0",
                     layoutInCell="1", allowOverlap="1").items():
        anchor.set(k, v)
    sp = etree.SubElement(anchor, f"{{{wp}}}simplePos")
    sp.set("x", "0")
    sp.set("y", "0")
    for tag in ("positionH", "positionV"):
        pos = etree.SubElement(anchor, f"{{{wp}}}{tag}")
        pos.set("relativeFrom", "page")
        off = etree.SubElement(pos, f"{{{wp}}}posOffset")
        off.text = "0"
    for hijo in list(inline):
        nombre = etree.QName(hijo).localname
        if nombre == "extent":
            anchor.append(hijo)
            ee = etree.SubElement(anchor, f"{{{wp}}}effectExtent")
            for a in "ltrb":
                ee.set(a, "0")
            etree.SubElement(anchor, f"{{{wp}}}wrapNone")
        elif nombre in ("docPr", "cNvGraphicFramePr", "graphic"):
            anchor.append(hijo)
    inline.getparent().replace(inline, anchor)


def _armar_encabezado(doc, texto_membrete):
    seccion = doc.sections[0]
    seccion.header.is_linked_to_previous = False
    enc = seccion.header

    # Párrafo 1: solo lleva la imagen de fondo (ocupa casi nada de alto)
    p_img = enc.paragraphs[0]
    _formato_parrafo(p_img, interlineado=1)
    if os.path.isfile(RUTA_MEMBRETE):
        _fondo_en_encabezado(doc, p_img, RUTA_MEMBRETE)
    else:
        print(f"[AVISO] No se encontró el membrete en: {RUTA_MEMBRETE}")

    # Párrafo 2: direcciones, debajo de "Alcaldía del Municipio Bolivariano de Brión"
    lineas = _lineas_membrete(texto_membrete)
    if not lineas:
        return
    largo = len(lineas) > 1
    p = enc.add_paragraph()
    _formato_parrafo(p, align=WD_ALIGN_PARAGRAPH.LEFT, antes=68, interlineado=7.2)
    p.paragraph_format.left_indent = Inches(2.42)
    p.paragraph_format.right_indent = Inches(1.4)
    run = p.add_run()
    _fuente(run, 6 if largo else 7, False, "1F1F1F", TIPO_LETRA_MEMBRETE)
    for i, linea in enumerate(lineas):
        run.add_text(linea)
        if i < len(lineas) - 1:
            run.add_break()


# ---------------------------------------------------------------------------
# FUNCIÓN PRINCIPAL
# ---------------------------------------------------------------------------
# Llamada desde app.py: generar_docx_ficha(data, fotos_paths, docx_path)
def generar_docx_ficha(data, fotos_paths, docx_path):
    doc = Document()

    # Corrección menor del archivo base de python-docx (Word/validadores exigen este atributo)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")

    # Página carta y márgenes (iguales a los del PDF)
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = Pt(MARGEN_LATERAL_PT)
    sec.right_margin = Pt(MARGEN_LATERAL_PT)
    sec.top_margin = Pt(MARGEN_SUPERIOR_PT)
    sec.bottom_margin = Pt(MARGEN_INFERIOR_PT)
    sec.header_distance = Pt(0)
    sec.footer_distance = Pt(0)

    # Estilo base y idioma (español, para que Word no subraye todo en rojo)
    normal = doc.styles["Normal"]
    normal.font.name = TIPO_LETRA
    normal.font.size = Pt(8.5)
    rpr = normal.element.get_or_add_rPr()
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "es-VE")
    lang.set(qn("w:eastAsia"), "es-VE")
    rpr.append(lang)

    doc.core_properties.title = "Ficha Técnica y Reporte de Avance de Obra"
    doc.core_properties.author = "Alcaldía del Municipio Bolivariano de Brión"

    # Membrete: fondo + direcciones elegidas en el menú
    _armar_encabezado(doc, data.get("membrete", ""))

    def titulo_seccion(texto, salto_antes=False, antes=8):
        p = doc.add_paragraph()
        _formato_parrafo(p, antes=antes, despues=4, interlineado=12,
                         mantener_con_siguiente=True, salto_antes=salto_antes)
        _escribir(p, texto, 10, True, AZUL)
        return p

    # --- TÍTULO (solo, en su misma posición) ---
    p = doc.add_paragraph()
    _formato_parrafo(p, align=WD_ALIGN_PARAGRAPH.CENTER, despues=8, interlineado=13)
    _escribir(p, "FICHA TÉCNICA Y REPORTE DE AVANCE DE OBRA", 10.5, True, AZUL)

    # --- I. DATOS GENERALES ---
    titulo_seccion("I. DATOS GENERALES", antes=0)

    columna_izq = [
        ("Proyecto:", data.get("proyecto", "")),
        ("Parroquia:", data.get("parroquia", "")),
        ("Ubicación:", data.get("ubicacion", "")),
        ("Ente Ejecutor:", data.get("ente_ejecutor", "")),
        ("Ente Inspector:", data.get("ente_inspector", "")),
        ("Tipología:", data.get("tipologia", "")),
        ("Tecnología Constructiva:", data.get("tecnologia", "")),
    ]
    columna_der = [
        ("Propietario:", data.get("propietario", "")),
        ("Cédula / RIF:", data.get("cedula", "")),
        ("Teléfono:", data.get("telefono", "")),
        ("Datos del Inmueble:", data.get("datos_inmueble", "")),
        ("Fecha de Emisión:", data.get("fecha_reporte", "")),
        ("Período Inspeccionado:", data.get("periodo_reporte", "")),
        ("Cantidad de Trabajadores:", data.get("cantidad_trabajadores", "")),
    ]

    t1 = _nueva_tabla(doc, len(columna_izq), [90, 175, 115, 150])
    for i, (izq, der) in enumerate(zip(columna_izq, columna_der)):
        fila = t1.rows[i]
        _fila_sin_partir(fila)
        for j, (etiqueta, valor) in enumerate((izq, der)):
            c_etq, c_val = fila.cells[j * 2], fila.cells[j * 2 + 1]
            _sombrear_celda(c_etq, GRIS_ETIQUETA)
            _texto_en_celda(c_etq, etiqueta, 8.5, True, AZUL)
            _texto_en_celda(c_val, valor or "", 8.5)
            c_etq.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            c_val.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    # --- II. TABLA DE AVANCE FÍSICO ---
    titulo_seccion("II. EVALUACIÓN Y TABLA DE AVANCE FÍSICO", antes=10)

    fases = data.get("fases", [])
    t2 = _nueva_tabla(doc, len(fases) + 1, [180, 50, 80, 80, 90], pad_v=2.5)
    encabezados = ["Capítulo / Fase Constructiva", "Peso (%)", "Avance Real (%)",
                   "Ponderado (%)", "Observación"]
    for j, texto in enumerate(encabezados):
        celda = t2.rows[0].cells[j]
        _sombrear_celda(celda, AZUL)
        _texto_en_celda(celda, texto, 8.5, True, "FFFFFF",
                        align=None if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
        celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    _fila_sin_partir(t2.rows[0])

    for i, fase in enumerate(fases, start=1):
        fila = t2.rows[i]
        _fila_sin_partir(fila)
        valores = [
            fase.get("nombre", ""),
            f"{fase.get('peso', 0)}%",
            f"{fase.get('avance_real', 0):.2f}%",
            f"{fase.get('ponderado', 0):.2f}%",
            fase.get("observacion", "") or "-",
        ]
        for j, valor in enumerate(valores):
            celda = fila.cells[j]
            _sombrear_celda(celda, "FFFFFF" if i % 2 == 1 else GRIS_FILA)
            _texto_en_celda(celda, valor, 8.5,
                            align=None if j == 0 else WD_ALIGN_PARAGRAPH.CENTER)
            celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    total = data.get("total_acumulado", 0)
    p = doc.add_paragraph()
    _formato_parrafo(p, antes=5, despues=8, interlineado=12)
    _escribir(p, f"PORCENTAJE TOTAL ACUMULADO DE LA OBRA: {total:.2f}%", 9, True, AZUL)

    # --- III. OBSERVACIONES GENERALES ---
    titulo_seccion("III. OBSERVACIONES GENERALES", antes=0)
    observaciones = (data.get("observaciones") or "").strip()
    p = doc.add_paragraph()
    _formato_parrafo(p, despues=10, interlineado=11)
    _escribir(p, observaciones or "Sin observaciones.", 8.5)

    # --- IV. REGISTRO FOTOGRÁFICO DE CAMPO ---
    # Con fotos, arranca en hoja nueva para que el título no se encime con ellas.
    titulo_seccion("IV. REGISTRO FOTOGRÁFICO DE CAMPO",
                   salto_antes=bool(fotos_paths), antes=0 if fotos_paths else 8)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    if not fotos_paths:
        p = doc.add_paragraph()
        _formato_parrafo(p, interlineado=11)
        _escribir(p, "No se registraron fotografías para este período.", 8.5)
    else:
        # Fotos cuadradas de 8 x 8 cm, 2 columnas (4 por hoja).
        lado = Cm(8.0)
        filas = (len(fotos_paths) + 1) // 2
        t3 = _nueva_tabla(doc, filas, [253.7, 253.7], borde_exterior=False,
                          rejilla=False, pad_v=10, pad_h=2)
        for idx, ruta in enumerate(fotos_paths):
            celda = t3.rows[idx // 2].cells[idx % 2]
            celda.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = celda.paragraphs[0]
            _formato_parrafo(p, align=WD_ALIGN_PARAGRAPH.CENTER)
            try:
                if not os.path.exists(ruta):
                    raise FileNotFoundError(f"no existe en disco: {ruta}")
                ruta_cuadrada = recortar_cuadrado(ruta)
                p.add_run().add_picture(ruta_cuadrada, width=lado, height=lado)
            except Exception as e:
                print(f"[AVISO] No se pudo cargar la foto {ruta}: {e}")
                _escribir(p, f"(imagen no disponible: {os.path.basename(str(ruta))})", 8.5)
        for fila in t3.rows:
            _fila_sin_partir(fila)

    doc.save(docx_path)
    return docx_path
